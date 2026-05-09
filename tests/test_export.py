from __future__ import annotations

import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from bb.category import CategorizedTransaction
from bb.config import BudgetBucketConfig, MonthlyBudgetConfig
from bb.export import export_yearly_budget_workbook
from bb.transaction import Transaction


class ExportTests(unittest.TestCase):
    def test_export_yearly_budget_workbook_writes_month_sheets(self) -> None:
        monthly_budget = MonthlyBudgetConfig(
            year=2026,
            month="march",
            needs=BudgetBucketConfig(frac=0.5, categories=("groceries",)),
            wants=BudgetBucketConfig(frac=0.3, categories=("movies",)),
            savings=BudgetBucketConfig(frac=0.2, categories=("savings",)),
        )

        march_transactions = [
            CategorizedTransaction(
                transaction=Transaction(
                    account="debit_demo",
                    source_file=Path("data/raw/debit_demo/2026/march.csv"),
                    posting_date="03/01/2026",
                    description="PAYROLL",
                    amount=Decimal("1000.00"),
                    source_metadata={},
                ),
                bucket="income",
                category="paycheck",
                matched_rule_id="paycheck_income",
                matched_rule=None,
                matched_on=(),
            ),
            CategorizedTransaction(
                transaction=Transaction(
                    account="debit_demo",
                    source_file=Path("data/raw/debit_demo/2026/march.csv"),
                    posting_date="03/03/2026",
                    description="Groceries",
                    amount=Decimal("-400.00"),
                    source_metadata={},
                ),
                bucket="needs",
                category="groceries",
                matched_rule_id="groceries",
                matched_rule=None,
                matched_on=(),
            ),
            CategorizedTransaction(
                transaction=Transaction(
                    account="credit_demo",
                    source_file=Path("data/raw/credit_demo/2026/march.csv"),
                    posting_date="03/05/2026",
                    description="Movie",
                    amount=Decimal("-250.00"),
                    source_metadata={},
                ),
                bucket="wants",
                category="movies",
                matched_rule_id="movies",
                matched_rule=None,
                matched_on=(),
            ),
        ]

        with tempfile.TemporaryDirectory() as tmpdir:
            destination = export_yearly_budget_workbook(
                lambda month: march_transactions if month == 3 else [],
                monthly_budget,
                Path(tmpdir),
            )

            self.assertEqual(destination, Path(tmpdir) / "2026.xlsx")
            self.assertTrue(destination.exists())

            workbook = load_workbook(destination)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "January",
                    "February",
                    "March",
                    "April",
                    "May",
                    "June",
                    "July",
                    "August",
                    "September",
                    "October",
                    "November",
                    "December",
                ],
            )
            worksheet = workbook["March"]

            values = {
                cell_value
                for row in worksheet.iter_rows(values_only=True)
                for cell_value in row
                if cell_value is not None
            }

            self.assertIn("March 2026", values)
            self.assertIn("Needs", values)
            self.assertIn("Categories", values)
            self.assertIn("$400.00", values)
            self.assertIn("40.0%", values)
            self.assertIn("Need Statistics", values)

    def test_export_yearly_budget_workbook_uses_file_path_when_suffix_is_present(self) -> None:
        monthly_budget = MonthlyBudgetConfig(year=2026, month="march")

        with tempfile.TemporaryDirectory() as tmpdir:
            destination = export_yearly_budget_workbook(
                lambda month: [],
                monthly_budget,
                Path(tmpdir) / "budget.csv",
            )

            self.assertEqual(destination, Path(tmpdir) / "budget.xlsx")
            self.assertTrue(destination.exists())


if __name__ == "__main__":
    unittest.main()

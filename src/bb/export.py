from __future__ import annotations

import calendar
from dataclasses import dataclass, replace
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from bb.category import CategorizedTransaction
from bb.config import MonthlyBudgetConfig
from bb.format import format_currency
from bb.report import (
    BudgetBucketPercentageComparison,
    compute_budget_target_vs_actual_percentages,
    compute_total_monthly_income,
    compute_total_non_income_inflows,
    compute_total_spend_by_category_within_bucket,
    format_category_label,
)

EXPORT_START_COLUMN = 2
EXPORT_STATS_START_COLUMN = 6
EXPORT_MONTH_TITLE_END_COLUMN = 6
EXPORT_INCOME_LABEL_COLUMN = 2
EXPORT_OTHER_INFLOWS_LABEL_COLUMN = 4


@dataclass(frozen=True)
class BudgetExportSection:
    heading: str
    category_rows: list[tuple[str, Decimal, Decimal | None]]
    stat_heading: str
    stat_rows: list[tuple[str, Decimal | str]]


def export_yearly_budget_workbook(
    categorized_transactions_for_year: callable,
    monthly_budget: MonthlyBudgetConfig,
    export_path: Path,
) -> Path:
    workbook = Workbook()
    for month in range(1, 13):
        month_budget = replace(monthly_budget, month=month)
        worksheet = workbook.active if month == 1 else workbook.create_sheet()
        worksheet.title = format_month_sheet_name(month)
        categorized_transactions = categorized_transactions_for_year(month)
        write_month_section(
            worksheet,
            categorized_transactions,
            month_budget,
            start_row=1,
        )
        set_workbook_column_widths(worksheet)

    destination = resolve_export_destination(export_path, monthly_budget.year)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination


def resolve_export_destination(export_path: Path, year: int) -> Path:
    if export_path.suffix:
        return export_path.with_suffix(".xlsx")
    return export_path / f"{year}.xlsx"


def build_month_sections(
    categorized_transactions: list[CategorizedTransaction],
    monthly_budget: MonthlyBudgetConfig,
) -> list[BudgetExportSection]:
    comparisons = compute_budget_target_vs_actual_percentages(
        categorized_transactions,
        monthly_budget,
    )
    spend_by_category = compute_total_spend_by_category_within_bucket(
        categorized_transactions
    )

    sections: list[BudgetExportSection] = []
    for bucket in ("needs", "wants", "savings"):
        comparison = comparisons[bucket]
        sections.append(
            BudgetExportSection(
                heading=format_category_label(bucket),
                category_rows=build_category_rows(
                    spend_by_category.get(bucket, {}),
                    comparison.monthly_income,
                )
                if bucket != "savings"
                else [],
                stat_heading=get_stat_heading(bucket),
                stat_rows=build_stat_rows(
                    comparison,
                    total_label="Total + Inflows" if bucket == "savings" else "Total",
                ),
            )
        )
    return sections


def build_category_rows(
    spend_by_category: dict[str, Decimal],
    monthly_income: Decimal,
) -> list[tuple[str, Decimal, Decimal | None]]:
    rows: list[tuple[str, Decimal, Decimal | None]] = []
    for category, total in sorted(spend_by_category.items()):
        percent = total / monthly_income if monthly_income > 0 else None
        rows.append((format_category_label(category), total, percent))
    return rows


def get_stat_heading(bucket: str) -> str:
    if bucket == "needs":
        return "Need Statistics"
    if bucket == "wants":
        return "Want Statistics"
    if bucket == "savings":
        return "Saving Statistics"
    return f"{format_category_label(bucket)} Statistics"


def build_stat_rows(
    comparison: BudgetBucketPercentageComparison,
    *,
    total_label: str,
) -> list[tuple[str, Decimal | str]]:
    actual_percent = (
        format_percentage_value(comparison.actual_fraction)
        if comparison.actual_fraction is not None
        else "n/a"
    )
    return [
        ("Target", comparison.monthly_income * comparison.target_fraction),
        (total_label, comparison.actual_spend),
        ("Target percent", format_percentage_value(comparison.target_fraction)),
        ("Percent of income", actual_percent),
    ]


def write_month_section(
    worksheet,
    categorized_transactions: list[CategorizedTransaction],
    monthly_budget: MonthlyBudgetConfig,
    *,
    start_row: int,
) -> int:
    month_label = format_month_heading(monthly_budget.year, monthly_budget.month)
    sections = build_month_sections(categorized_transactions, monthly_budget)
    write_month_heading(worksheet, start_row, month_label)
    write_month_financials(worksheet, categorized_transactions, start_row + 1)
    base_row = start_row + 4

    needs_end_row = write_section_tables(
        worksheet,
        sections[0],
        start_row=base_row,
        category_start_column=EXPORT_START_COLUMN,
        stats_start_column=EXPORT_STATS_START_COLUMN,
    )
    wants_end_row = write_section_tables(
        worksheet,
        sections[1],
        start_row=needs_end_row + 2,
        category_start_column=EXPORT_START_COLUMN,
        stats_start_column=EXPORT_STATS_START_COLUMN,
    )
    savings_start_row = wants_end_row + 2
    savings_end_row = write_section_tables(
        worksheet,
        sections[2],
        start_row=savings_start_row,
        category_start_column=EXPORT_START_COLUMN,
        stats_start_column=EXPORT_STATS_START_COLUMN,
    )

    return savings_end_row


def write_month_heading(worksheet, row: int, month_label: str) -> None:
    worksheet.merge_cells(
        start_row=row,
        start_column=EXPORT_START_COLUMN,
        end_row=row,
        end_column=EXPORT_MONTH_TITLE_END_COLUMN,
    )
    cell = worksheet.cell(row=row, column=EXPORT_START_COLUMN, value=month_label)
    cell.font = Font(bold=True, size=14)
    cell.alignment = Alignment(horizontal="center")


def write_month_financials(
    worksheet,
    categorized_transactions: list[CategorizedTransaction],
    row: int,
) -> None:
    income = compute_total_monthly_income(categorized_transactions)
    other_inflows = compute_total_non_income_inflows(categorized_transactions)
    entries = [
        (row, EXPORT_INCOME_LABEL_COLUMN, "Income", format_currency(income)),
        (
            row + 1,
            EXPORT_INCOME_LABEL_COLUMN,
            "Other inflows",
            format_currency(other_inflows),
        ),
    ]
    for target_row, start_column, label, value in entries:
        label_cell = worksheet.cell(row=target_row, column=start_column, value=label)
        label_cell.font = Font(bold=True)
        label_cell.alignment = Alignment(horizontal="center")
        value_cell = worksheet.cell(
            row=target_row,
            column=start_column + 1,
            value=value,
        )
        value_cell.alignment = Alignment(horizontal="center")


def write_section_tables(
    worksheet,
    section: BudgetExportSection,
    *,
    start_row: int,
    category_start_column: int,
    stats_start_column: int,
) -> int:
    category_end_row = start_row
    if section.category_rows:
        category_end_row = write_category_table(
            worksheet,
            section.heading,
            section.category_rows,
            start_row=start_row,
            start_column=category_start_column,
        )

    stat_end_row = write_stat_table(
        worksheet,
        section.stat_heading,
        section.stat_rows,
        start_row=start_row,
        start_column=stats_start_column,
    )
    return max(category_end_row, stat_end_row)


def write_category_table(
    worksheet,
    heading: str,
    rows: list[tuple[str, Decimal, Decimal | None]],
    *,
    start_row: int,
    start_column: int,
) -> int:
    headers = ["Categories", "Values", "Percent"]
    return write_table(
        worksheet,
        headers=headers,
        rows=[
            [label, format_currency(value), format_optional_percentage(percent)]
            for label, value, percent in rows
        ],
        start_row=start_row,
        start_column=start_column,
        title=heading,
    )


def write_stat_table(
    worksheet,
    heading: str,
    rows: list[tuple[str, Decimal | str]],
    *,
    start_row: int,
    start_column: int,
) -> int:
    return write_table(
        worksheet,
        headers=[heading, "Value"],
        rows=[
            [label, format_currency(value) if isinstance(value, Decimal) else value]
            for label, value in rows
        ],
        start_row=start_row,
        start_column=start_column,
    )


def write_table(
    worksheet,
    *,
    headers: list[str],
    rows: list[list[str]],
    start_row: int,
    start_column: int,
    title: str | None = None,
) -> int:
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    title_fill = PatternFill("solid", fgColor="EDEDED")

    if title is not None:
        worksheet.merge_cells(
            start_row=start_row,
            start_column=start_column,
            end_row=start_row,
            end_column=start_column + len(headers) - 1,
        )
        title_cell = worksheet.cell(row=start_row, column=start_column, value=title)
        title_cell.font = Font(bold=True)
        title_cell.alignment = Alignment(horizontal="center")
        title_cell.fill = title_fill
        for column in range(start_column, start_column + len(headers)):
            worksheet.cell(row=start_row, column=column).border = border
        start_row += 1

    for offset, header in enumerate(headers):
        cell = worksheet.cell(row=start_row, column=start_column + offset, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = header_fill
        cell.border = border

    for row_offset, row_values in enumerate(rows, start=1):
        for column_offset, value in enumerate(row_values):
            cell = worksheet.cell(
                row=start_row + row_offset,
                column=start_column + column_offset,
                value=value,
            )
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

    return start_row + len(rows)


def set_workbook_column_widths(worksheet) -> None:
    uniform_width = 18
    widths = {index: uniform_width for index in range(1, 8)}
    for column_index, width in widths.items():
        worksheet.column_dimensions[get_column_letter(column_index)].width = width


def format_optional_percentage(fraction: Decimal | None) -> str:
    if fraction is None:
        return "n/a"
    return format_percentage_value(fraction)


def format_percentage_value(fraction: Decimal) -> str:
    return f"{(fraction * Decimal('100')):.1f}%"


def format_month_sheet_name(month: int) -> str:
    return calendar.month_name[month]


def format_month_heading(year: int, month: str | int) -> str:
    if isinstance(month, int):
        month_number = month
    else:
        month_number = (
            int(month)
            if str(month).isdigit()
            else list(calendar.month_name).index(str(month).title())
        )
    return f"{calendar.month_name[month_number]} {year}"
